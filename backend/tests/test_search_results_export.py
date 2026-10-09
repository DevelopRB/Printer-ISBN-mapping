"""Tests for ISBN Search Results Excel export — Print Date mapping."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from services.search_results_export import SEARCH_EXPORT_HEADERS, build_search_results_excel


def _sheet_rows(xlsx_bytes: bytes) -> list[list[object]]:
    wb = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb.active
    return [[cell for cell in row] for row in ws.iter_rows(values_only=True)]


def test_headers_include_print_date():
    assert "Print Date" in SEARCH_EXPORT_HEADERS
    assert SEARCH_EXPORT_HEADERS.index("Print Date") == 3


def test_print_date_exported_for_found_isbn():
    """Same field the UI uses: print_date from the search occurrence."""
    data = build_search_results_excel(
        [
            {
                "isbn": "9788182054370",
                "list_a": "N/A",
                "printer": "Printer 1",
                "print_date": "2026-08-24",  # UI field: occurrence.print_date
                "job_end": "10:00",
                "print_start": "09:50",
                "job_accept": "09:40",
                "result": "Complete",
                "set_number": "2",
                "output_set": "2",
                "source_filename": "9788182054370.pdf",
                "status": "FOUND",
            }
        ]
    )
    rows = _sheet_rows(data)
    assert rows[0][3] == "Print Date"
    assert rows[1][0] == "9788182054370"
    assert rows[1][3] == "24-Aug-2026"


def test_multiple_occurrences_keep_individual_print_dates():
    data = build_search_results_excel(
        [
            {
                "isbn": "9788182054370",
                "list_a": "N/A",
                "printer": "Printer 1",
                "print_date": "2026-08-24",
                "status": "FOUND",
            },
            {
                "isbn": "9788182054370",
                "list_a": "N/A",
                "printer": "Printer 1",
                "print_date": "2026-08-26",
                "status": "FOUND",
            },
            {
                "isbn": "9788182054370",
                "list_a": "N/A",
                "printer": "Printer 2",
                "print_date": "2026-08-27",
                "status": "FOUND",
            },
        ]
    )
    rows = _sheet_rows(data)
    dates = [r[3] for r in rows[1:]]
    assert dates == ["24-Aug-2026", "26-Aug-2026", "27-Aug-2026"]


def test_not_found_has_blank_print_date():
    data = build_search_results_excel(
        [
            {
                "isbn": "9789353242497",
                "list_a": "N/A",
                "printer": "",
                "print_date": "",
                "status": "NOT FOUND",
            }
        ]
    )
    rows = _sheet_rows(data)
    assert rows[1][0] == "9789353242497"
    assert rows[1][3] in (None, "")
    assert rows[1][11] == "NOT FOUND"


def test_multiple_isbns_export_dates():
    data = build_search_results_excel(
        [
            {
                "isbn": "9788182054370",
                "print_date": "2026-08-24",
                "printer": "Printer 1",
                "status": "FOUND",
            },
            {
                "isbn": "9789362800398",
                "print_date": "2026-08-20",
                "printer": "Printer 1",
                "status": "FOUND",
            },
            {
                "isbn": "8888001171897",
                "print_date": "2026-08-27",
                "printer": "Printer 2",
                "status": "FOUND",
            },
        ]
    )
    rows = _sheet_rows(data)
    by_isbn = {r[0]: r[3] for r in rows[1:]}
    assert by_isbn["9788182054370"] == "24-Aug-2026"
    assert by_isbn["9789362800398"] == "20-Aug-2026"
    assert by_isbn["8888001171897"] == "27-Aug-2026"
