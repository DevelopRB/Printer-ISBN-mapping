"""Search-data must use the same printer ISBN set as normalized Excel export."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from database import SessionLocal, init_db
from models.models import AutomatedRun
from services.normalized_export import classify_run_downloads, export_normalized_for_run


@pytest.fixture(scope="module")
def completed_run_id():
    init_db()
    db = SessionLocal()
    try:
        run = (
            db.query(AutomatedRun)
            .filter(AutomatedRun.status.in_(("COMPLETED", "PARTIAL_SUCCESS")))
            .order_by(AutomatedRun.id.desc())
            .first()
        )
        if not run:
            pytest.skip("No completed automated run in local DB")
        return run.id
    finally:
        db.close()


def test_search_data_matches_tsv_classify_and_excel(completed_run_id, client=None):
    """Printer occurrence ISBNs in search-data == classify_run_downloads == Excel Sheet 2."""
    from fastapi.testclient import TestClient
    from main import app

    db = SessionLocal()
    try:
        classified = classify_run_downloads(db, completed_run_id)
        assert classified is not None
        tsv_isbns = [t.normalized_identifier for t in classified.extracted if t.normalized_identifier]

        with TestClient(app) as tc:
            resp = tc.get(f"/api/reconciliation/run/{completed_run_id}/search-data")
            assert resp.status_code == 200
            payload = resp.json()
            search_isbns = [o["isbn"] for o in payload["printer_occurrences"]]

        assert search_isbns == tsv_isbns, (
            "Search must return the same ordered printer ISBN occurrences as Excel classify path"
        )

        buf = export_normalized_for_run(db, completed_run_id)
        wb = load_workbook(buf, read_only=True, data_only=True)
        excel_isbns: list[str] = []
        for name in wb.sheetnames:
            if not name.startswith("Printer"):
                continue
            ws = wb[name]
            rows = ws.iter_rows(min_row=2, values_only=True)
            for row in rows:
                if row and row[0]:
                    excel_isbns.append(str(row[0]))
        wb.close()

        assert sorted(search_isbns) == sorted(excel_isbns)
        assert len(search_isbns) == len(excel_isbns)
    finally:
        db.close()


def test_known_isbn_from_wider_tsv_extractable():
    """Sanity: extract_isbn still returns the reported ISBN from its filename."""
    from services.isbn_extractor import extract_isbn

    assert extract_isbn("9788182054370_Pharma..pdf") == "9788182054370"
