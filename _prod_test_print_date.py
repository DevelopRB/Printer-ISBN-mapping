"""Live validation: Print Date from same job row for Printers 1–3."""

from __future__ import annotations

import json
import time
import uuid
from pathlib import Path

import urllib.error
import urllib.request
from openpyxl import load_workbook

BASE = "http://127.0.0.1:8000/api"
LIST_A = Path(r"c:\Users\abhis\Downloads\pending.xlsx")
OUT = Path(r"c:\Users\abhis\Desktop\Printer ISBN mapping\_prod_test_out")
OUT.mkdir(exist_ok=True)


def req(method, path, data=None, files=None, timeout=600):
    url = BASE + path
    if files:
        boundary = uuid.uuid4().hex
        body = b""
        for name, (fname, content, ctype) in files.items():
            body += f"--{boundary}\r\n".encode()
            body += (
                f'Content-Disposition: form-data; name="{name}"; '
                f'filename="{fname}"\r\n'
            ).encode()
            body += f"Content-Type: {ctype}\r\n\r\n".encode()
            body += content + b"\r\n"
        body += f"--{boundary}--\r\n".encode()
        request = urllib.request.Request(
            url, data=body, headers={"Content-Type": f"multipart/form-data; boundary={boundary}"}, method=method
        )
    elif data is not None:
        request = urllib.request.Request(
            url,
            data=json.dumps(data).encode(),
            headers={"Content-Type": "application/json"},
            method=method,
        )
    else:
        request = urllib.request.Request(url, method=method)
    try:
        with urllib.request.urlopen(request, timeout=timeout) as resp:
            raw = resp.read()
            if raw[:2] == b"PK":
                return raw
            return json.loads(raw.decode())
    except urllib.error.HTTPError as exc:
        raise RuntimeError(f"{exc.code}: {exc.read().decode(errors='replace')}") from exc


def wait_run(run_id):
    for i in range(180):
        time.sleep(2)
        r = req("GET", f"/reconciliation/run/{run_id}")
        print(f"  poll {i+1}: {r['status']}")
        if r["status"] in ("COMPLETED", "PARTIAL_SUCCESS", "FAILED"):
            return r
    raise TimeoutError(run_id)


def find_header(ws, *titles):
    for r in range(1, 15):
        vals = [ws.cell(r, c).value for c in range(1, len(titles) + 1)]
        if vals == list(titles):
            return r
    raise AssertionError(f"header {titles} not found")


def main():
    # Probe all three printers for date+filename same-row
    from datetime import date
    from database import SessionLocal
    from printers.joblog_adapter import build_joblog_adapter_from_config
    from services.printer_config_service import ensure_network_printer_configs

    db = SessionLocal()
    configs = {c.slot: c for c in ensure_network_printer_configs(db)}
    probe_results = {}
    for slot in (1, 2, 3):
        adapter = build_joblog_adapter_from_config(configs[slot])
        try:
            res = adapter.fetch_export(date(2026, 8, 25), date(2026, 8, 27))
            sample = [
                {
                    "job_id": r.job_no,
                    "print_date": r.printer_date,
                    "filename": r.source_filename,
                }
                for r in res.records[:3]
            ]
            with_date = sum(1 for r in res.records if r.printer_date)
            probe_results[slot] = {
                "ok": res.ok,
                "count": res.record_count,
                "with_print_date": with_date,
                "sample": sample,
            }
            print(f"Printer {slot}: ok={res.ok} count={res.record_count} dated={with_date}")
            for s in sample:
                print(" ", s)
        finally:
            adapter.disconnect()
    db.close()
    (OUT / "print_date_probe.json").write_text(json.dumps(probe_results, indent=2), encoding="utf-8")

    # Full run Printer 1
    la = req(
        "POST",
        "/reconciliation/list-a/upload",
        files={
            "file": (
                LIST_A.name,
                LIST_A.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    print("List A", la["id"], la["isbn_detected_count"])
    run = req(
        "POST",
        "/reconciliation/run",
        data={
            "list_a_id": la["id"],
            "printer_ids": [1],
            "start_date": "2026-08-01",
            "end_date": "2026-08-26",
        },
    )
    run = wait_run(run["id"])
    print("run", run["id"], run["status"], run["downloads"][0]["record_count"])
    assert run["status"] in ("COMPLETED", "PARTIAL_SUCCESS")

    wb_bytes = req("GET", f"/reconciliation/run/{run['id']}/export/normalized")
    out = OUT / "print_date_normalized.xlsx"
    out.write_bytes(wb_bytes)
    book = load_workbook(out, data_only=True)

    # Meta on sheet 1
    ws1 = book["List A ISBNs"]
    meta = {ws1.cell(r, 1).value: ws1.cell(r, 2).value for r in range(1, 12) if ws1.cell(r, 1).value}
    print("META", {k: meta[k] for k in meta if k in ("Printer", "Report Date Range", "Downloaded On", "ISBN Records")})
    assert "01-Aug-2026" in (meta.get("Report Date Range") or "")
    assert "26-Aug-2026" in (meta.get("Report Date Range") or "")

    ws2 = book["Printer ISBNs"]
    hdr = find_header(ws2, "ISBN", "Print Date", "Source Filename")
    rows = []
    for r in range(hdr + 1, min(hdr + 6, ws2.max_row + 1)):
        rows.append([ws2.cell(r, c).value for c in range(1, 4)])
    print("Sheet2 sample", rows)
    # Every ISBN must appear in its filename; print date must be ops format not download date alone
    checked = 0
    for r in range(hdr + 1, ws2.max_row + 1):
        isbn = str(ws2.cell(r, 1).value or "")
        pdate = str(ws2.cell(r, 2).value or "")
        src = str(ws2.cell(r, 3).value or "")
        if not isbn:
            continue
        digits = "".join(ch for ch in src if ch.isdigit())
        assert isbn in digits, (isbn, src)
        assert "-Aug-" in pdate or "-Jul-" in pdate or pdate == "", pdate
        checked += 1
        if checked >= 50:
            break
    print("validated first", checked, "sheet2 rows")

    ws3 = book["Reconciliation"]
    hdr3 = find_header(ws3, "List A ISBN", "Printer ISBN", "Print Date", "Status")
    present = 0
    for r in range(hdr3 + 1, ws3.max_row + 1):
        if ws3.cell(r, 4).value == "PRESENT":
            present += 1
            assert ws3.cell(r, 3).value, f"PRESENT without print date row {r}"
    print("PRESENT rows with print dates", present)
    print("PASS")


if __name__ == "__main__":
    main()
