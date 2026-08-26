"""Ordered occurrence validation: manual printer 4010 vs automated Sheet 2.

Compares ORDERED LISTS (stronger than Counter / unique sets).

Usage:
  python scripts/compare_printer_identifier_sets.py ^
    --manual "C:\\Users\\abhis\\Desktop\\text.xlsx" ^
    --normalized "C:\\Users\\abhis\\Downloads\\Normalized_Reconciliation_....xlsx"

  python scripts/compare_printer_identifier_sets.py ^
    --manual "C:\\Users\\abhis\\Desktop\\text.xlsx" ^
    --tsv "backend\\data\\printer_downloads\\printer_1_....tsv"
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from services.isbn_extractor import extract_isbn, split_4010_filenames
from services.normalized_export import (
    extract_traces_from_4010_values,
    identifier_traceable_to_raw,
    ordered_occurrence_lists_match,
    printer_isbns_from_tsv_bytes,
)


def load_manual_4010(path: Path) -> list[str]:
    """Load 4010 values from a printer export (TSV or TSV saved with .xlsx name)."""
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return []

    header = lines[0]
    if "\t" in header:
        parts0 = header.split("\t")
        try:
            idx = next(
                i
                for i, h in enumerate(parts0)
                if h.strip() in ("4010", "Filename", "filename")
            )
        except StopIteration:
            idx = 1 if len(parts0) > 1 else 0
        rows = []
        for line in lines[1:]:
            parts = line.split("\t")
            if idx < len(parts) and parts[idx].strip():
                rows.append(parts[idx].strip().strip('"'))
        return rows

    if header.strip().lower() in ("4010", "filename"):
        return [ln.strip().strip('"') for ln in lines[1:] if ln.strip()]

    raise ValueError(f"Unsupported manual file format: {path}")


def extract_all(raws: list[str]) -> list[str]:
    """Ordered occurrence list (source order; duplicates preserved)."""
    return [t.normalized_identifier for t in extract_traces_from_4010_values(raws)]


def load_sheet2(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Printer ISBNs"]
    return [
        str(ws.cell(r, 1).value).strip()
        for r in range(2, ws.max_row + 1)
        if ws.cell(r, 1).value
    ]


def report_ordered(label_a: str, a: list[str], label_b: str, b: list[str]) -> None:
    ca, cb = Counter(a), Counter(b)
    ordered_ok = ordered_occurrence_lists_match(a, b)
    print(f"{label_a} occurrences / unique:", len(a), len(ca))
    print(f"{label_b} occurrences / unique:", len(b), len(cb))
    print("ORDERED OCCURRENCE MATCH:", ordered_ok)
    print("len equal:", len(a) == len(b))
    print("Counter equal:", ca == cb)
    print("Unique set equal:", set(a) == set(b))
    if not ordered_ok:
        # First mismatch index
        n = min(len(a), len(b))
        mismatch_i = next((i for i in range(n) if a[i] != b[i]), None)
        if len(a) != len(b):
            print(f"Length mismatch: {label_a}={len(a)} {label_b}={len(b)}")
        if mismatch_i is not None:
            print(
                f"First order mismatch at i={mismatch_i}: "
                f"{label_a}={a[mismatch_i]!r} {label_b}={b[mismatch_i]!r}"
            )
        only_a = ca - cb
        only_b = cb - ca
        print(f"{label_a}-only (occurrence deltas):", dict(only_a) if only_a else 0)
        print(f"{label_b}-only (occurrence deltas):", dict(only_b) if only_b else 0)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--manual", required=True)
    ap.add_argument("--normalized", default=None)
    ap.add_argument("--tsv", default=None, help="Current-run printer TSV to compare")
    args = ap.parse_args()

    raws = load_manual_4010(Path(args.manual))
    no_isbn = 0
    for raw in raws:
        parts = split_4010_filenames(raw) or [raw]
        if not any(extract_isbn(fn) for fn in parts):
            no_isbn += 1
    manual_occurrences = extract_all(raws)

    print("Manual 4010 rows:", len(raws))
    print("Manual records without extractable ISBN:", no_isbn)
    print("Manual ISBN occurrences:", len(manual_occurrences))
    print()

    if args.tsv:
        traces = printer_isbns_from_tsv_bytes(
            Path(args.tsv).read_bytes(), source_label="tsv"
        )
        automated_occurrences = [t.normalized_identifier for t in traces]
        print("--- Manual vs TSV extract (ORDERED) ---")
        report_ordered("Manual", manual_occurrences, "TSV", automated_occurrences)
        print()
        untraced = [
            t.normalized_identifier
            for t in traces
            if not identifier_traceable_to_raw(t.normalized_identifier, t.raw_4010_value)
        ]
        print("TSV traces failing provenance assert:", len(untraced))
        # Spot-check: occurrence_index is sequential
        assert [t.occurrence_index for t in traces] == list(range(1, len(traces) + 1))

    if args.normalized:
        automated_occurrences = load_sheet2(Path(args.normalized))
        print("--- Manual vs Normalized Sheet 2 (ORDERED) ---")
        report_ordered("Manual", manual_occurrences, "Sheet2", automated_occurrences)
        print()
        untraced = [
            isbn
            for isbn in set(automated_occurrences)
            if not any(identifier_traceable_to_raw(isbn, raw) for raw in raws)
        ]
        print("Sheet 2 ISBNs with no provenance in manual 4010:", len(untraced))
        if untraced:
            print(untraced)


if __name__ == "__main__":
    main()
