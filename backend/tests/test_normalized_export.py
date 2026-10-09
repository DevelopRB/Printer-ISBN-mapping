"""Tests for the Operations normalized export workbook."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from services.normalized_export import (
    PRESENT,
    NOT_FOUND,
    SKIP_REASON_NO_ISBN,
    build_normalized_workbook,
)


def _sheets(buffer):
    buffer.seek(0)
    wb = openpyxl.load_workbook(buffer, data_only=True)
    return wb


def test_sheet_names_and_order_legacy_single_sheet():
    buf = build_normalized_workbook(
        ["8888001171897", "9789372259803"],
        [("8888001171897", "8888001171897_F2.pdf")],
        [(2, "CE Adjust", SKIP_REASON_NO_ISBN)],
    )
    wb = _sheets(buf)
    assert wb.sheetnames == [
        "List A ISBNs",
        "Printer ISBNs",
        "Reconciliation",
        "Skipped ISBNs",
    ]


def test_multi_printer_separate_sheets():
    buf = build_normalized_workbook(
        ["8888001171897", "1111012883952"],
        skipped_rows=[
            (2, "CE Adjust", SKIP_REASON_NO_ISBN, "Printer 1"),
            (5, "Details Counter", SKIP_REASON_NO_ISBN, "Printer 2"),
        ],
        printer_sheets=[
            (
                "Printer 1 ISBNs",
                [("8888001171897", "2026-08-26", "8888001171897_F2.pdf")],
            ),
            (
                "Printer 2 ISBNs",
                [
                    ("8888001171897", "2026-08-27", "8888001171897_F3.pdf"),
                    ("1111012883952", "2026-08-27", "1111012883952_F2.pdf"),
                ],
            ),
            (
                "Printer 3 ISBNs",
                [("9789372259803", "2026-08-25", "9789372259803_F2.pdf")],
            ),
        ],
    )
    wb = _sheets(buf)
    assert wb.sheetnames == [
        "List A ISBNs",
        "Printer 1 ISBNs",
        "Printer 2 ISBNs",
        "Printer 3 ISBNs",
        "Reconciliation",
        "Skipped ISBNs",
    ]
    assert wb["Printer 1 ISBNs"].cell(2, 1).value == "8888001171897"
    assert wb["Printer 2 ISBNs"].max_row == 3
    assert wb["Printer 3 ISBNs"].cell(2, 1).value == "9789372259803"
    # Reconciliation uses combined presence across printers
    recon = wb["Reconciliation"]
    assert recon.cell(1, 10).value == "Printer"
    assert recon.cell(1, 11).value == "Status"
    assert recon.cell(2, 11).value == PRESENT  # 8888… from P1/P2
    assert recon.cell(3, 11).value == PRESENT  # 1111… from P2
    # Print dates aggregated across printers for same ISBN
    assert "26-Aug-2026" in (recon.cell(2, 3).value or "")
    assert "27-Aug-2026" in (recon.cell(2, 3).value or "")
    assert recon.cell(2, 10).value == "Printer 1, Printer 2"
    assert recon.cell(3, 10).value == "Printer 2"
    skip = wb["Skipped ISBNs"]
    assert skip.cell(1, 1).value == "Printer"
    assert skip.cell(2, 1).value == "Printer 1"


def test_list_a_preserves_duplicates():
    buf = build_normalized_workbook(
        ["8888001171897", "8888001171897", "9789372259803"],
        [("8888001171897", "a.pdf")],
    )
    ws = _sheets(buf)["List A ISBNs"]
    values = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    assert values == ["8888001171897", "8888001171897", "9789372259803"]


def test_printer_preserves_duplicates_and_filenames():
    rows = [
        ("8888001171897", "2026-08-26", "8888001171897_F2.pdf"),
        ("8888001171897", "2026-08-27", "8888001171897_F3.pdf"),
        ("8888001171897", "2026-08-27", "8888001171897_F4.pdf"),
        ("9789372259803", "2026-08-26", "9789372259803_F2.pdf"),
    ]
    buf = build_normalized_workbook(["8888001171897"], rows)
    ws = _sheets(buf)["Printer ISBNs"]
    assert ws.cell(1, 1).value == "ISBN"
    assert ws.cell(1, 2).value == "Print Date"
    assert ws.cell(1, 11).value == "Source Filename"
    values = [
        (ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 11).value)
        for r in range(2, ws.max_row + 1)
    ]
    assert values == [
        ("8888001171897", "26-Aug-2026", "8888001171897_F2.pdf"),
        ("8888001171897", "27-Aug-2026", "8888001171897_F3.pdf"),
        ("8888001171897", "27-Aug-2026", "8888001171897_F4.pdf"),
        ("9789372259803", "26-Aug-2026", "9789372259803_F2.pdf"),
    ]


def test_reconciliation_preserves_list_a_duplicates():
    list_a = [
        "8888001171897",
        "8888001171897",
        "1111012883952",
        "9789372259803",
        "9789375480896",
    ]
    printer = [
        ("8888001171897", "2026-08-26", "a.pdf"),
        ("8888001171897", "2026-08-27", "b.pdf"),
        ("1111012883952", "2026-08-26", "c.pdf"),
        ("9789372259803", "2026-08-26", "d.pdf"),
    ]
    buf = build_normalized_workbook(list_a, printer)
    ws = _sheets(buf)["Reconciliation"]
    rows = [
        (
            ws.cell(r, 1).value,
            ws.cell(r, 2).value or "",
            ws.cell(r, 3).value or "",
            ws.cell(r, 11).value,
        )
        for r in range(2, ws.max_row + 1)
    ]
    # Same count/order as List A; printer duplicates do not add extra reconciliation rows
    assert rows == [
        ("8888001171897", "8888001171897", "26-Aug-2026, 27-Aug-2026", PRESENT),
        ("8888001171897", "8888001171897", "26-Aug-2026, 27-Aug-2026", PRESENT),
        ("1111012883952", "1111012883952", "26-Aug-2026", PRESENT),
        ("9789372259803", "9789372259803", "26-Aug-2026", PRESENT),
        ("9789375480896", "", "", NOT_FOUND),
    ]


def test_printer_only_not_in_sheet3():
    buf = build_normalized_workbook(
        ["8888001171897"],
        [("8888001171897", "a.pdf"), ("1111012883952", "b.pdf")],
    )
    ws = _sheets(buf)["Reconciliation"]
    assert ws.max_row == 2  # header + 1 data row
    assert ws.cell(2, 1).value == "8888001171897"
    assert ws.cell(2, 11).value == PRESENT


def test_business_identifier_not_converted():
    buf = build_normalized_workbook(
        ["8888000240648"],
        [("8888000240648", "8888000240648_F2.pdf")],
    )
    ws = _sheets(buf)["Reconciliation"]
    assert ws.cell(2, 1).value == "8888000240648"
    assert ws.cell(2, 2).value == "8888000240648"
    assert ws.cell(2, 11).value == PRESENT
    assert not str(ws.cell(2, 1).value).startswith("978888")


def test_isbn_stored_as_text_not_scientific():
    buf = build_normalized_workbook(
        ["9788121255905"],
        [("9788121255905", "9788121255905_F2.pdf")],
    )
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)
    for sheet in ("List A ISBNs", "Printer ISBNs", "Reconciliation"):
        cell = wb[sheet].cell(2, 1)
        assert cell.value == "9788121255905"
        assert isinstance(cell.value, str)
        assert cell.number_format == "@"


def test_hyphenated_4010_is_traceable_provenance():
    from services.normalized_export import identifier_traceable_to_raw

    raw = "Pub_978-93-6433-007-7 F2 _ PB Book A New and Comprehensive_Pharma..pdf"
    assert identifier_traceable_to_raw("9789364330077", raw) is True
    assert identifier_traceable_to_raw("9789364330077", "8888001171897_F2.pdf") is False


def test_provenance_rejects_unrelated_isbn():
    from services.normalized_export import identifier_traceable_to_raw

    assert identifier_traceable_to_raw("9789372259803", "8888000240648_F2_Pharma.pdf") is False


def test_extract_traces_skips_non_isbn_records_preserves_dupes():
    from services.normalized_export import extract_traces_from_4010_values

    raws = [
        "CE Adjust",
        "8888001171897_F2.pdf",
        "8888001171897_F3.pdf",
        "Ethiopian Bible Old Testament_merge_Pharma..pdf",
        "8888001171897_F4.pdf",
    ]
    traces = extract_traces_from_4010_values(raws)
    ids = [t.normalized_identifier for t in traces]
    assert ids == ["8888001171897", "8888001171897", "8888001171897"]
    assert [t.occurrence_index for t in traces] == [1, 2, 3]
    assert all(t.raw_4010_value for t in traces)


def test_skipped_isbns_sheet_preserves_source_order():
    from services.normalized_export import classify_4010_values

    raws = [
        "8888003685861_F2_QC LB 20 COPY_Pharma.pdf",  # 1
        "CE Adjust",  # 2 skipped
        "CE Adjust",  # 3 skipped
        "8888000446965_F2_QC LB 20 COPY_Pharma.pdf",  # 4
        "Details Counter",  # 5 skipped
        "8888000446965_F2_QC LB 20 COPY_Pharma.pdf",  # 6
        "VERSION 2 MAMTA_Pharma.pdf",  # 7 skipped
        "Pub_9788182054608-F2-HB-BOOK-(2-COPY)-DEEPTI_Pharma.pdf",  # 8
    ]
    classified = classify_4010_values(raws)
    assert [t.normalized_identifier for t in classified.extracted] == [
        "8888003685861",
        "8888000446965",
        "8888000446965",
        "9788182054608",
    ]
    assert [(s.source_row, s.printer_record, s.reason) for s in classified.skipped] == [
        (2, "CE Adjust", SKIP_REASON_NO_ISBN),
        (3, "CE Adjust", SKIP_REASON_NO_ISBN),
        (5, "Details Counter", SKIP_REASON_NO_ISBN),
        (7, "VERSION 2 MAMTA_Pharma.pdf", SKIP_REASON_NO_ISBN),
    ]
    assert classified.raw_record_count == 8

    printer_rows = [
        (t.normalized_identifier, t.document_filename or "")
        for t in classified.extracted
    ]
    skipped_rows = [(s.source_row, s.printer_record, s.reason) for s in classified.skipped]
    buf = build_normalized_workbook(["8888003685861"], printer_rows, skipped_rows)
    ws = _sheets(buf)["Skipped ISBNs"]
    assert [ws.cell(1, c).value for c in range(1, 4)] == [
        "Source Row",
        "Printer Record / Filename",
        "Reason",
    ]
    sheet_rows = [
        (ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value)
        for r in range(2, ws.max_row + 1)
    ]
    assert sheet_rows == [
        ("2", "CE Adjust", SKIP_REASON_NO_ISBN),
        ("3", "CE Adjust", SKIP_REASON_NO_ISBN),
        ("5", "Details Counter", SKIP_REASON_NO_ISBN),
        ("7", "VERSION 2 MAMTA_Pharma.pdf", SKIP_REASON_NO_ISBN),
    ]


def test_sheet2_preserves_source_order_with_gaps_and_duplicates():
    """CE Adjust gaps must not reorder later ISBNs; duplicates stay adjacent."""
    from services.normalized_export import (
        extract_traces_from_4010_values,
        ordered_occurrence_lists_match,
        build_normalized_workbook,
    )

    raws = [
        "8888003685861_F2_QC LB 20 COPY_Pharma.pdf",
        "CE Adjust",
        "8888000446965_F2_QC LB 20 COPY_Pharma.pdf",
        "8888000446965_F2_QC LB 20 COPY_Pharma.pdf",
        "VERSION 2 MAMTA_Pharma.pdf",
        "Pub_9788182054608-F2-HB-BOOK-(2-COPY)-DEEPTI_Pharma.pdf",
    ]
    expected = [
        "8888003685861",
        "8888000446965",
        "8888000446965",
        "9788182054608",
    ]
    traces = extract_traces_from_4010_values(raws)
    occurrences = [t.normalized_identifier for t in traces]
    assert occurrences == expected
    assert ordered_occurrence_lists_match(expected, occurrences)

    wrong_order = [
        "8888003685861",
        "8888000446965",
        "9788182054608",
        "8888000446965",
    ]
    assert not ordered_occurrence_lists_match(expected, wrong_order)

    printer_rows = [(t.normalized_identifier, t.document_filename or "") for t in traces]
    buf = build_normalized_workbook(["8888003685861"], printer_rows)
    ws = _sheets(buf)["Printer ISBNs"]
    values = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    assert values == expected


def test_pipe_split_preserves_both_occurrences_in_order():
    from services.normalized_export import classify_4010_values

    raw = (
        "9789374567555_F3_QC_LB_Pharma.pdf | "
        "9789374567555_F3_LB_QC_Pharma.pdf"
    )
    classified = classify_4010_values([raw])
    ids = [t.normalized_identifier for t in classified.extracted]
    assert ids == ["9789374567555", "9789374567555"]
    assert classified.skipped == []
    assert classified.raw_record_count == 1


def test_sheet2_from_tsv_matches_4010_pipeline():
    from services.normalized_export import (
        classify_4010_values,
        classify_tsv_bytes,
        ordered_occurrence_lists_match,
    )

    tsv = (
        "0\t1122\t4001\t4009\t4010\t4005\t1630\t1034\n"
        "1\t2026\t8\t1\t8888001171897_F2.pdf\t\t\t\n"
        "2\t2026\t8\t1\tCE Adjust\t\t\t\n"
        "3\t2026\t8\t1\t8888001171897_F3.pdf\t\t\t\n"
    ).encode("utf-8")
    classified = classify_tsv_bytes(tsv)
    ids = [t.normalized_identifier for t in classified.extracted]
    assert ids == ["8888001171897", "8888001171897"]
    assert [(s.source_row, s.printer_record) for s in classified.skipped] == [
        (2, "CE Adjust")
    ]
    raws = ["8888001171897_F2.pdf", "CE Adjust", "8888001171897_F3.pdf"]
    from_raw = classify_4010_values(raws)
    assert ordered_occurrence_lists_match(
        [t.normalized_identifier for t in from_raw.extracted],
        ids,
    )
